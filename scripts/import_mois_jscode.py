#!/usr/bin/env python3
"""Normalize MOIS jscode Excel files (KiKcd_H / KiKcd_B / KiKmix).

The official archive contains both fixed-width TEXT and XLSX. XLSX is preferred
for normalization; the raw archive remains the audit source.
"""
from __future__ import annotations

import argparse, csv, json, re
from pathlib import Path
from typing import Any
from openpyxl import load_workbook


def s(v: Any) -> str:
    if v is None: return ""
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v).strip()


def date8(v: Any) -> str:
    x = re.sub(r"\D", "", s(v))
    return f"{x[:4]}-{x[4:6]}-{x[6:8]}" if len(x) == 8 else ""


def status_from_end(v: Any) -> str:
    return "ABOLISHED" if date8(v) else "CURRENT"


def find_xlsx(root: Path, prefix: str) -> Path:
    matches = sorted(p for p in root.rglob("*.xlsx") if p.name.lower().startswith(prefix.lower()))
    if not matches: raise FileNotFoundError(prefix)
    return matches[0]


def read_xlsx(path: Path) -> tuple[list[str], list[dict[str,str]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    headers = [s(v) for v in next(it)]
    rows=[]
    for raw in it:
        vals=[s(v) for v in raw]
        if not any(vals): continue
        vals += [""] * (len(headers)-len(vals))
        rows.append(dict(zip(headers, vals)))
    return headers, rows


def join_name(*parts: str) -> str:
    return " ".join(p.strip() for p in parts if p and p.strip())


def legal_type(code: str, emd: str) -> tuple[int,str]:
    if code[2:] == "00000000": return 1, "SIDO"
    if code[5:] == "00000": return 2, "SIGUNGU"
    if code[8:] == "00":
        if emd.endswith("읍"): return 3, "LEGAL_EUP"
        if emd.endswith("면"): return 3, "LEGAL_MYEON"
        return 3, "LEGAL_DONG"
    return 4, "LEGAL_RI"


def admin_type(code: str, emd: str) -> tuple[int,str]:
    if code[2:] == "00000000": return 1, "ADMIN_SIDO"
    if code[5:] == "00000": return 2, "ADMIN_SIGUNGU"
    if emd.endswith("읍"): return 3, "ADMIN_EUP"
    if emd.endswith("면"): return 3, "ADMIN_MYEON"
    return 3, "ADMIN_DONG"


def parent_code(code: str, level: int) -> str:
    if level == 1: return ""
    if level == 2: return code[:2] + "00000000"
    if level == 3: return code[:5] + "00000"
    if level == 4: return code[:8] + "00"
    raise ValueError(level)


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("extracted_dir", type=Path)
    ap.add_argument("--snapshot-date", required=True)
    ap.add_argument("--out-dir", type=Path, default=Path("data/normalized"))
    a=ap.parse_args(); a.out_dir.mkdir(parents=True, exist_ok=True)

    hp=find_xlsx(a.extracted_dir,"KIKcd_H"); bp=find_xlsx(a.extracted_dir,"KIKcd_B"); mp=find_xlsx(a.extracted_dir,"KIKmix")
    hh,hr=read_xlsx(hp); bh,br=read_xlsx(bp); mh,mr=read_xlsx(mp)

    admin=[]
    for r in hr:
        code=re.sub(r"\D","",r.get("행정동코드","")); sido=r.get("시도명",""); sgg=r.get("시군구명",""); emd=r.get("읍면동명","")
        if len(code)!=10: continue
        level,ptype=admin_type(code,emd); full=join_name(sido,sgg,emd); pc=parent_code(code,level)
        admin.append({"place_id":f"adm:{code}","place_type":ptype,"hierarchy_level":level,"name_ko":(emd or sgg or sido),"full_name_ko":full,"official_code":code,"parent_place_id":f"adm:{pc}" if pc else "","legal_status":status_from_end(r.get("말소일자")),"valid_from":date8(r.get("생성일자")),"valid_to":date8(r.get("말소일자")),"source_id":"mois_jscode","source_snapshot_date":a.snapshot_date})

    legal=[]
    for r in br:
        code=re.sub(r"\D","",r.get("법정동코드","")); sido=r.get("시도명",""); sgg=r.get("시군구명",""); emd=r.get("읍면동명",""); ri=r.get("동리명","")
        if len(code)!=10: continue
        level,ptype=legal_type(code,emd); full=join_name(sido,sgg,emd,ri); pc=parent_code(code,level)
        legal.append({"place_id":f"bjd:{code}","place_type":ptype,"hierarchy_level":level,"name_ko":(ri or emd or sgg or sido),"full_name_ko":full,"official_code":code,"parent_place_id":f"bjd:{pc}" if pc else "","legal_status":status_from_end(r.get("말소일자")),"valid_from":date8(r.get("생성일자")),"valid_to":date8(r.get("말소일자")),"source_id":"mois_jscode","source_snapshot_date":a.snapshot_date})

    relations=[]
    for r in mr:
        ac=re.sub(r"\D","",r.get("행정동코드","")); lc=re.sub(r"\D","",r.get("법정동코드",""))
        if len(ac)!=10 or len(lc)!=10: continue
        relations.append({"from_place_id":f"adm:{ac}","to_place_id":f"bjd:{lc}","relation_type":"ADMINISTERS","legal_status":status_from_end(r.get("말소일자")),"valid_from":date8(r.get("생성일자")),"valid_to":date8(r.get("말소일자")),"source_id":"mois_jscode","source_snapshot_date":a.snapshot_date})

    def write(name, rows):
        p=a.out_dir/name
        with p.open("w",encoding="utf-8-sig",newline="") as f:
            w=csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)
    write("places_admin.csv",admin); write("places_legal_mois.csv",legal); write("relations_admin_legal.csv",relations)

    meta={"snapshot_date":a.snapshot_date,"files":{"admin":hp.name,"legal":bp.name,"mapping":mp.name},"headers":{"admin":hh,"legal":bh,"mapping":mh},"rows":{"admin":len(admin),"legal":len(legal),"mapping":len(relations)},"current":{"admin":sum(x["legal_status"]=="CURRENT" for x in admin),"legal":sum(x["legal_status"]=="CURRENT" for x in legal),"mapping":sum(x["legal_status"]=="CURRENT" for x in relations)},"abolished":{"admin":sum(x["legal_status"]=="ABOLISHED" for x in admin),"legal":sum(x["legal_status"]=="ABOLISHED" for x in legal),"mapping":sum(x["legal_status"]=="ABOLISHED" for x in relations)}}
    (a.out_dir/"mois_jscode_import_meta.json").write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(meta,ensure_ascii=False,indent=2))

if __name__=="__main__": main()
