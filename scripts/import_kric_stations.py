#!/usr/bin/env python3
from __future__ import annotations
import argparse,csv,hashlib,json,re
from collections import Counter,defaultdict
from pathlib import Path
from openpyxl import load_workbook

def s(v):
    if v is None:return ''
    return str(v).strip()
def norm(v):return re.sub(r'\s+',' ',s(v)).strip()
def keynorm(v):return re.sub(r'[^0-9A-Za-z가-힣]','',norm(v)).lower()
def pick(headers,*cands):
    m={keynorm(h):h for h in headers}
    for c in cands:
        if keynorm(c) in m:return m[keynorm(c)]
    return None
def fnum(v):
    try:return float(str(v).strip())
    except:return None
def read_csv(path:Path):
    if not path.exists():return []
    with path.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def write(path,rows,fields):
    path.parent.mkdir(parents=True,exist_ok=True)
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore');w.writeheader();w.writerows(rows)

SIDO_PREFIX={
 '서울시':'서울특별시','서울':'서울특별시','부산시':'부산광역시','부산':'부산광역시','대구시':'대구광역시','대구':'대구광역시',
 '인천시':'인천광역시','인천':'인천광역시','광주시':'광주광역시','광주':'광주광역시','대전시':'대전광역시','대전':'대전광역시',
 '울산시':'울산광역시','울산':'울산광역시','세종시':'세종특별자치시','세종':'세종특별자치시','경기':'경기도',
 '강원도':'강원특별자치도','강원':'강원특별자치도','충북':'충청북도','충남':'충청남도','전북':'전북특별자치도','전남':'전라남도',
 '경북':'경상북도','경남':'경상남도','제주':'제주특별자치도','제주도':'제주특별자치도'}

def valid_at(r,date):
    vf=(r.get('valid_from') or '').strip(); vt=(r.get('valid_to') or '').strip()
    return (not vf or vf<=date) and (not vt or date<vt or date==vt)

def correction_rules(path:Path):return read_csv(path) if path else []
def normalize_address(address,station,operator,lines,rules):
    raw=norm(address); a=raw; method=[]
    for r in rules:
        if r.get('station_name')==station and a.startswith(r.get('raw_prefix','')):
            a=r.get('corrected_prefix','')+a[len(r.get('raw_prefix','')):];method.append('AUDITED_CORRECTION');break
    tok=a.split(' ',1)[0] if a else ''
    if tok in SIDO_PREFIX:
        a=SIDO_PREFIX[tok]+a[len(tok):]; method.append('SIDO_ALIAS')
    if ('대구교통공사' in operator or '대구 도시철도' in lines) and re.match(r'^(중구|동구|서구|남구|북구|수성구|달서구|달성군)',a):
        a='대구광역시 '+a;method.append('OPERATOR_JURISDICTION')
    if '구리도시공사' in operator and not re.match(r'^(경기도|구리시)',a):
        a='경기도 구리시 '+a;method.append('OPERATOR_JURISDICTION')
    return raw,norm(a),method

def main():
    ap=argparse.ArgumentParser();ap.add_argument('xlsx',type=Path);ap.add_argument('--master-places',type=Path,required=True);ap.add_argument('--rules-dir',type=Path,default=Path('data/rules'));ap.add_argument('--out-dir',type=Path,default=Path('data/normalized/stations'));ap.add_argument('--snapshot-date',default='2026-06-30');a=ap.parse_args()
    wb=load_workbook(a.xlsx,read_only=True,data_only=True);ws=wb[wb.sheetnames[0]];it=ws.iter_rows(values_only=True);headers=[s(x) for x in next(it)]
    cols={'station_no':pick(headers,'역번호','역사번호'),'name':pick(headers,'역사명','역명'),'line_no':pick(headers,'노선번호'),'line':pick(headers,'노선명','선명'),'name_en':pick(headers,'영문역사명','영문역명','영문명'),'name_hanja':pick(headers,'한자역사명','한자역명'),'transfer':pick(headers,'환승역구분','환승역여부'),'transfer_lines':pick(headers,'환승노선명','환승노선'),'transfer_line_numbers':pick(headers,'환승노선번호'),'lat':pick(headers,'역위도','역사위도','위도'),'lon':pick(headers,'역경도','역사경도','경도'),'operator':pick(headers,'운영기관명','철도운영기관명'),'address':pick(headers,'역사도로명주소','도로명주소','주소'),'phone':pick(headers,'역사전화번호','전화번호'),'date':pick(headers,'데이터기준일자')}
    if not cols['name']:raise SystemExit(f'역사명 컬럼 없음: {headers}')
    raw=[]
    for vals in it:
        r=dict(zip(headers,[s(x) for x in vals]));name=norm(r.get(cols['name'],''))
        if name:raw.append({k:(norm(r.get(c,'')) if c else '') for k,c in cols.items()})
    groups=defaultdict(list)
    for r in raw:
        addr=keynorm(r['address']);lat=fnum(r['lat']);lon=fnum(r['lon'])
        key=f"{keynorm(r['name'])}|{addr}" if addr else (f"{keynorm(r['name'])}|{lat:.4f}|{lon:.4f}" if lat is not None and lon is not None else f"{keynorm(r['name'])}|{keynorm(r['operator'])}|{keynorm(r['line'])}")
        groups[key].append(r)
    master=read_csv(a.master_places);regions=[r for r in master if r.get('place_type') in {'SIDO','SIGUNGU'} and valid_at(r,a.snapshot_date)];regions.sort(key=lambda r:(r.get('place_type')=='SIGUNGU',len(r.get('full_name_ko',''))),reverse=True)
    sigungu=[r for r in regions if r.get('place_type')=='SIGUNGU'];rules=correction_rules(a.rules_dir/'station_address_corrections.csv')
    places=[];lines=[];aliases=[];rels=[];unmatched=[];methods=Counter()
    for k,items in groups.items():
        first=items[0];pid=f"station:kric:{hashlib.sha1(k.encode()).hexdigest()[:16]}";lats=[fnum(x['lat']) for x in items if fnum(x['lat']) is not None];lons=[fnum(x['lon']) for x in items if fnum(x['lon']) is not None];lat=sum(lats)/len(lats) if lats else None;lon=sum(lons)/len(lons) if lons else None
        address=next((x['address'] for x in items if x['address']),'');operators='|'.join(sorted({x['operator'] for x in items if x['operator']}));line_names='|'.join(sorted({x['line'] for x in items if x['line']}))
        places.append({'place_id':pid,'place_type':'STATION','hierarchy_level':'','name_ko':first['name'],'full_name_ko':first['name']+'역','official_code':'','parent_place_id':'','legal_status':'CURRENT','valid_from':'','valid_to':'','source_id':'kric_urban_station','validity_source_id':'kric_urban_station','source_snapshot_date':a.snapshot_date,'latitude':f'{lat:.7f}' if lat is not None else '','longitude':f'{lon:.7f}' if lon is not None else ''})
        aliases += [{'place_id':pid,'alias':first['name'],'alias_type':'STATION_NAME','source_id':'kric_urban_station'},{'place_id':pid,'alias':first['name']+'역','alias_type':'STATION_NAME_WITH_SUFFIX','source_id':'kric_urban_station'}]
        for x in items:
            if x['name_en']:aliases.append({'place_id':pid,'alias':x['name_en'],'alias_type':'ENGLISH_NAME','source_id':'kric_urban_station'})
            if x['name_hanja']:aliases.append({'place_id':pid,'alias':x['name_hanja'],'alias_type':'HANJA_NAME','source_id':'kric_urban_station'})
            lines.append({'station_place_id':pid,'station_number':x['station_no'],'line_number':x['line_no'],'line_name':x['line'],'operator_name':x['operator'],'is_transfer':x['transfer'],'transfer_lines':x['transfer_lines'],'transfer_line_numbers':x['transfer_line_numbers'],'road_address':x['address'],'phone':x['phone'],'data_date':x['date'] or a.snapshot_date,'source_id':'kric_urban_station'})
        raw_addr,na,steps=normalize_address(address,first['name'],operators,line_names,rules);matched=None;method=''
        if na:
            for rr in regions:
                full=norm(rr.get('full_name_ko',''))
                if full and (na==full or na.startswith(full+' ') or full in na):matched=rr;method='FULL_NAME';break
            if not matched:
                cands={r['place_id']:r for r in sigungu if na.startswith(norm(r.get('name_ko',''))+' ') or re.search(r'(^| )'+re.escape(norm(r.get('name_ko','')))+r'( |$)',na)}
                if len(cands)==1:matched=next(iter(cands.values()));method='UNIQUE_SIGUNGU'
        if matched:
            confidence='1.0000' if method=='FULL_NAME' and not steps else ('0.9900' if 'AUDITED_CORRECTION' in steps or 'SIDO_ALIAS' in steps else '0.9500')
            if 'OPERATOR_JURISDICTION' in steps:confidence='0.9000'
            methods['+'.join(steps+[method]) if steps else method]+=1
            rels.append({'from_place_id':pid,'to_place_id':matched['place_id'],'relation_type':'STATION_IN','confidence':confidence,'source_id':'kric_urban_station','valid_from':'','valid_to':''})
        else:unmatched.append({'station_place_id':pid,'station_name':first['name'],'address':raw_addr,'normalized_address':na,'lines':line_names,'operators':operators})
    seen=set();dedup=[]
    for r in aliases:
        kk=(r['place_id'],r['alias'],r['alias_type'])
        if r['alias'] and kk not in seen:seen.add(kk);dedup.append(r)
    a.out_dir.mkdir(parents=True,exist_ok=True)
    write(a.out_dir/'places_station.csv',places,['place_id','place_type','hierarchy_level','name_ko','full_name_ko','official_code','parent_place_id','legal_status','valid_from','valid_to','source_id','validity_source_id','source_snapshot_date','latitude','longitude'])
    write(a.out_dir/'station_lines.csv',lines,['station_place_id','station_number','line_number','line_name','operator_name','is_transfer','transfer_lines','transfer_line_numbers','road_address','phone','data_date','source_id'])
    write(a.out_dir/'station_aliases.csv',dedup,['place_id','alias','alias_type','source_id']);write(a.out_dir/'station_relations.csv',rels,['from_place_id','to_place_id','relation_type','confidence','source_id','valid_from','valid_to']);write(a.out_dir/'station_unmatched_regions.csv',unmatched,['station_place_id','station_name','address','normalized_address','lines','operators'])
    report={'raw_rows':len(raw),'physical_stations':len(places),'stations_with_coordinates':sum(bool(r['latitude'] and r['longitude']) for r in places),'line_memberships':len(lines),'aliases':len(dedup),'region_relations':len(rels),'unmatched_regions':len(unmatched),'match_methods':dict(methods),'headers':headers,'mapped_columns':cols}
    (a.out_dir/'station_import_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
